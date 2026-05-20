# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

VENDOR_DIR = Path(__file__).resolve().parent / ".vendor"
if VENDOR_DIR.exists():
    sys.path.insert(0, str(VENDOR_DIR))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


REPO_ROOT = Path(__file__).resolve().parent
ASSETS_DIR = REPO_ROOT / "assets"
DOCS_DIR = REPO_ROOT / "docs"
OUTPUT_DIR = REPO_ROOT / "course_outputs"
PPTX_PATH = OUTPUT_DIR / "基于AlphaZero与蒙特卡洛树搜索的五子棋AI设计_课程汇报.pptx"

BG = RGBColor(8, 13, 20)
PANEL = RGBColor(19, 30, 44)
ACCENT = RGBColor(78, 161, 255)
ACCENT_2 = RGBColor(88, 214, 141)
TEXT = RGBColor(240, 244, 248)
SUBTLE = RGBColor(166, 178, 189)


def ensure_materials():
    OUTPUT_DIR.mkdir(exist_ok=True)
    summary_path = OUTPUT_DIR / "experiment_summary.json"
    if not summary_path.exists():
        subprocess.run([sys.executable, "run_course_experiments.py"], cwd=REPO_ROOT, check=True)
    subprocess.run([sys.executable, "generate_course_docs.py"], cwd=REPO_ROOT, check=True)


def load_summary():
    summary_path = OUTPUT_DIR / "experiment_summary.json"
    if summary_path.exists():
        return json.loads(summary_path.read_text(encoding="utf-8"))
    return {
        "baseline": {"final_loss": "-", "best_win_ratio": "-", "last_win_ratio": "-"},
        "reward_shaping": {
            "final_loss": "-",
            "best_win_ratio": "-",
            "last_win_ratio": "-",
            "avg_heuristic_reward": "-",
            "total_open_three": "-",
            "total_open_four": "-",
            "total_five": "-",
        },
        "comparison": {"loss_delta": "-", "win_ratio_delta": "-"},
    }


def set_run_font(run, size, bold=False, color=TEXT, name="Arial", east_asia="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    r_pr = run._r.get_or_add_rPr()
    r_pr.set(qn("a:ea"), east_asia)


def set_text_frame(text_frame, text, size=24, bold=False, color=TEXT, align=PP_ALIGN.LEFT, name="Arial"):
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color, name=name)


def add_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.20)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()


def add_title(slide, kicker, title, subtitle=None):
    add_background(slide)
    tx = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(12), Inches(1.0))
    set_text_frame(tx.text_frame, kicker, size=14, bold=True, color=ACCENT)
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.95), Inches(12), Inches(1.0))
    set_text_frame(title_box.text_frame, title, size=28, bold=True, color=TEXT)
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.55), Inches(12), Inches(0.6))
        set_text_frame(sub_box.text_frame, subtitle, size=13, color=SUBTLE)


def add_bullets(slide, bullets, left=0.8, top=2.0, width=5.5, height=4.5, font_size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = "• " + bullet
        p.level = 0
        p.space_after = Pt(8)
        for run in p.runs:
            set_run_font(run, size=font_size, color=TEXT, name="Arial")
    return box


def add_image(slide, path, left, top, width, height=None):
    if Path(path).exists():
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), None if height is None else Inches(height))
    else:
        placeholder = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height or 3.5)
        )
        placeholder.fill.solid()
        placeholder.fill.fore_color.rgb = PANEL
        placeholder.line.color.rgb = ACCENT
        set_text_frame(placeholder.text_frame, "Placeholder", size=18, color=SUBTLE, align=PP_ALIGN.CENTER)


def add_table(slide, rows, cols, data, left=0.7, top=2.0, width=12.0, height=3.0):
    table = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height)).table
    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if r > 0 else ACCENT
            cell.text = str(data[r][c])
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    set_run_font(run, size=14, bold=(r == 0), color=TEXT)
    return table


def footer(slide, page_no):
    box = slide.shapes.add_textbox(Inches(11.6), Inches(6.9), Inches(1.1), Inches(0.3))
    set_text_frame(box.text_frame, f"{page_no:02d}", size=12, color=SUBTLE, align=PP_ALIGN.RIGHT)


def build_presentation(summary):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_background(slide)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(9.8), Inches(1.8))
    set_text_frame(title_box.text_frame, "基于 AlphaZero 与蒙特卡洛树搜索的五子棋 AI 设计", size=30, bold=True)
    sub = slide.shapes.add_textbox(Inches(0.82), Inches(2.2), Inches(8.8), Inches(0.8))
    set_text_frame(sub.text_frame, "课程设计汇报 / AlphaZero_Gomoku on Windows + Conda + PyTorch", size=16, color=SUBTLE)
    meta = slide.shapes.add_textbox(Inches(0.82), Inches(5.65), Inches(4.8), Inches(1.0))
    set_text_frame(meta.text_frame, "姓名：________   学校：________   日期：2026-05-19", size=16, color=TEXT)
    add_image(slide, ASSETS_DIR / "architecture_diagram.png", 8.8, 1.2, 3.9, 4.6)
    footer(slide, 1)

    slides_spec = [
        ("背景", "研究背景", ["博弈 AI 从规则搜索发展到深度强化学习。", "AlphaGo 展示了策略网络、价值网络与 MCTS 的组合威力。", "AlphaZero 进一步取消人类棋谱监督，完全依赖自博弈。"], ASSETS_DIR / "alpha_zero_workflow.png"),
        ("目标", "项目目标", ["实现可运行的五子棋人机对战系统。", "实现基于自博弈的数据采样与策略学习。", "在 Windows + Conda 环境下稳定训练与实验复现。"], ASSETS_DIR / "module_relationship.png"),
        ("架构", "系统总体架构", ["Game Environment 提供状态转移与胜负判定。", "MCTS 负责带先验的搜索决策。", "Policy-Value Network 提供策略概率与局面价值。", "Training Pipeline 串联 self-play、训练、评估。"], ASSETS_DIR / "architecture_diagram.png"),
        ("原理", "AlphaZero 原理", ["自博弈产生训练样本。", "策略头输出动作先验。", "价值头输出局面胜率估计。", "MCTS 用网络引导搜索，再反过来监督网络。"], ASSETS_DIR / "alpha_zero_workflow.png"),
        ("搜索", "MCTS 原理", ["Selection：基于 Q+U 选择最优分支。", "Expansion：扩展叶子节点的可行动作。", "Simulation：用价值网络评估叶子局面。", "Backup：把叶子价值沿路径逐层回传。"], ASSETS_DIR / "mcts_tree.png"),
        ("网络", "神经网络结构", ["输入 4 个平面：己方、对手、上一步、先手标记。", "共享 3 层卷积抽取棋盘局部特征。", "Policy head 输出每个动作概率。", "Value head 输出 [-1,1] 的局面价值。"], ASSETS_DIR / "architecture_diagram.png"),
        ("流程", "Self-play 流程", ["MCTS 依据网络先验搜索动作。", "保存 state、pi、current_player。", "终局后构造 z；若启用 shaping，则构造 shaped_z。", "数据增强后加入 replay buffer。"], ASSETS_DIR / "alpha_zero_workflow.png"),
        ("训练", "训练流程", ["收集自博弈数据。", "当缓冲区足够大时执行 policy_update。", "每隔若干 batch 与纯 MCTS 对弈评估。", "保存 current / best policy。"], ASSETS_DIR / "experiment_loss.png"),
        ("代码", "项目代码结构", ["game.py：棋盘与自博弈", "mcts_alphaZero.py：搜索核心", "policy_value_net_pytorch.py：网络推理与训练", "train.py：主训练循环", "heuristic_reward.py：课程设计创新点"], ASSETS_DIR / "module_relationship.png"),
        ("创新", "课程设计创新点", ["选择方案 A：启发式奖励 shaping。", "以最小改动在 self-play 奖励端加入活三、活四、连五先验。", "通过参数开关保证可回退、可对比、可复现实验。"], ASSETS_DIR / "experiment_heuristic_reward.png"),
        ("实验", "实验设计", ["对比组：Baseline vs Reward shaping。", "棋盘：6x6，4 连胜。", "MCTS playout：32，batch size：32，batch 数：8。", "评估指标：loss、win ratio、启发式统计、推理时间。"], ASSETS_DIR / "experiment_win_ratio.png"),
        ("结果", "实验结果", ["Reward shaping 已稳定产生日志与模式计数。", f"Final loss: baseline={summary['baseline']['final_loss']} / shaping={summary['reward_shaping']['final_loss']}", f"Best win ratio: baseline={summary['baseline']['best_win_ratio']} / shaping={summary['reward_shaping']['best_win_ratio']}", "短程实验中胜率差异不大，但 shaping 注入了更强的局面结构信号。"], ASSETS_DIR / "experiment_loss.png"),
        ("效率", "推理时间与对战展示", ["使用提供的 8x8 预训练模型测试不同 playout 的单步推理时间。", "人机对战演示已在本地跑通，可作为课程展示截图。"], ASSETS_DIR / "inference_time.png"),
        ("展示", "AI 对战展示", ["当前项目已跑通 human_play.py。", "可以使用预训练模型进行命令行人机对战。", "课程答辩中可展示固定脚本化对局，也可现场手动输入坐标。"], REPO_ROOT / "artifacts" / "screenshots" / "human_play.png"),
        ("总结", "总结与展望", ["项目已完成 Windows + Conda + PyTorch 跑通。", "创新点采用 reward shaping，改动小且可稳定运行。", "后续可扩展到动态 c_puct、transposition table 或更完整棋形库。"], ASSETS_DIR / "module_relationship.png"),
    ]

    page = 2
    for kicker, title, bullets, image_path in slides_spec:
        slide = prs.slides.add_slide(blank)
        add_title(slide, kicker, title)
        add_bullets(slide, bullets)
        add_image(slide, image_path, 7.0, 1.8, 5.6, 4.6)
        footer(slide, page)
        page += 1

    slide = prs.slides.add_slide(blank)
    add_title(slide, "对比", "Baseline 与 Reward Shaping 对比表", "来自自动实验脚本生成的结果摘要")
    data = [
        ["指标", "Baseline", "Reward shaping"],
        ["Final loss", summary["baseline"]["final_loss"], summary["reward_shaping"]["final_loss"]],
        ["Best win ratio", summary["baseline"]["best_win_ratio"], summary["reward_shaping"]["best_win_ratio"]],
        ["Last win ratio", summary["baseline"]["last_win_ratio"], summary["reward_shaping"]["last_win_ratio"]],
        ["Avg heuristic reward", "0", summary["reward_shaping"]["avg_heuristic_reward"]],
        ["Open three count", "0", summary["reward_shaping"]["total_open_three"]],
        ["Open four count", "0", summary["reward_shaping"]["total_open_four"]],
        ["Five count", "0", summary["reward_shaping"]["total_five"]],
    ]
    add_table(slide, len(data), 3, data, top=1.9, height=3.8)
    add_image(slide, ASSETS_DIR / "experiment_win_ratio.png", 7.5, 4.2, 5.0, 2.6)
    footer(slide, page)

    prs.save(str(PPTX_PATH))


def main():
    ensure_materials()
    build_presentation(load_summary())
    print(PPTX_PATH)


if __name__ == "__main__":
    main()
